from __future__ import annotations

import logging
import os
import tempfile
from pathlib import Path

import fitz  # type: ignore[import-untyped]
from docx import Document

from src.ingest.models import DocumentBlock

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS: frozenset[str] = frozenset({".docx", ".pdf", ".jpg", ".jpeg", ".png", ".doc"})


def parse_file(file_path: str | Path) -> list[DocumentBlock]:
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext not in SUPPORTED_EXTENSIONS:
        logger.warning("Unsupported file extension: %s (file: %s)", ext, path.name)
        return []

    if ext == ".docx":
        return _parse_docx(path)
    elif ext == ".pdf":
        return _parse_pdf(path)
    elif ext in {".jpg", ".jpeg", ".png"}:
        return _parse_image(path)
    elif ext == ".doc":
        logger.warning(
            "Old .doc format requires LibreOffice headless conversion (file: %s). "
            "Use SOFFICE_PATH to enable conversion.",
            path.name,
        )
        return _parse_doc_via_soffice(path)

    return []


def _parse_docx(path: Path) -> list[DocumentBlock]:
    try:
        doc = Document(str(path))
    except Exception as e:
        logger.error("Failed to open docx: %s (%s)", path.name, e)
        return []

    blocks: list[DocumentBlock] = []
    current_chapter = ""
    block_index = 0

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = para.style.name.lower() if para.style else ""
        if style_name.startswith("heading"):
            current_chapter = text
        blocks.append(
            DocumentBlock(
                source_file=path.name,
                source_title=path.stem,
                content=text,
                doc_type="paragraph",
                chapter=current_chapter,
                block_index=block_index,
            )
        )
        block_index += 1

    for _table_idx, table in enumerate(doc.tables):
        for _row_idx, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(cells)
            if row_text.strip():
                blocks.append(
                    DocumentBlock(
                        source_file=path.name,
                        source_title=path.stem,
                        content=row_text,
                        doc_type="table",
                        chapter=current_chapter,
                        page_num=None,
                        block_index=block_index,
                    )
                )
                block_index += 1

    logger.info("Parsed docx: %s -> %d blocks", path.name, len(blocks))
    return blocks


def _parse_pdf(path: Path) -> list[DocumentBlock]:
    try:
        doc = fitz.open(str(path))
    except Exception as e:
        logger.error("Failed to open pdf: %s (%s)", path.name, e)
        return []

    blocks: list[DocumentBlock] = []
    block_index = 0
    full_text_length = 0

    for page_num in range(len(doc)):
        page = doc[page_num]
        page_text = page.get_text().strip()
        full_text_length += len(page_text)

        if page_text:
            blocks.append(
                DocumentBlock(
                    source_file=path.name,
                    source_title=path.stem,
                    content=page_text,
                    doc_type="pdf_text",
                    chapter="",
                    page_num=page_num + 1,
                    block_index=block_index,
                )
            )
            block_index += 1
        else:
            logger.info("PDF page %d has no extractable text, will need OCR", page_num + 1)

    doc.close()

    if full_text_length < 50 and blocks:
        logger.warning(
            "PDF %s has very little text (%d chars) - pages may be scanned images requiring OCR",
            path.name,
            full_text_length,
        )

    logger.info("Parsed pdf: %s -> %d blocks (%d chars)", path.name, len(blocks), full_text_length)
    return blocks


def _parse_image(path: Path) -> list[DocumentBlock]:
    logger.info("Image file %s - OCR will be needed for text extraction", path.name)
    return [
        DocumentBlock(
            source_file=path.name,
            source_title=path.stem,
            content=f"[Image file: {path.name} - process via GLM OCR for text extraction]",
            doc_type="image",
            chapter="",
            block_index=0,
        )
    ]


def _parse_doc_via_soffice(path: Path) -> list[DocumentBlock]:
    soffice_path = os.getenv("SOFFICE_PATH", "")
    if not soffice_path:
        logger.warning("SOFFICE_PATH not set, cannot convert .doc file: %s", path.name)
        return []

    import subprocess

    with tempfile.TemporaryDirectory() as tmpdir:
        output_path = os.path.join(tmpdir, f"{path.stem}.docx")
        try:
            subprocess.run(
                [soffice_path, "--headless", "--convert-to", "docx", "--outdir", tmpdir, str(path)],
                capture_output=True,
                text=True,
                timeout=60,
                check=True,
            )
        except subprocess.TimeoutExpired:
            logger.error("LibreOffice conversion timed out for: %s", path.name)
            return []
        except subprocess.CalledProcessError as e:
            logger.error("LibreOffice conversion failed for %s: %s", path.name, e.stderr)
            return []

        if os.path.isfile(output_path):
            return _parse_docx(Path(output_path))

    return []


def list_source_files(source_dir: str | Path) -> list[Path]:
    source_path = Path(source_dir)
    if not source_path.is_dir():
        logger.warning("Source directory not found: %s", source_dir)
        return []

    files: list[Path] = []
    for ext in SUPPORTED_EXTENSIONS:
        files.extend(source_path.rglob(f"*{ext}"))

    files.sort()

    logger.info("Found %d source files in %s", len(files), source_dir)
    return files
