from __future__ import annotations

import unittest
from io import BytesIO
from unittest.mock import patch

from docx import Document

from src.models import MaterialSlot
from src.services.field_extractor import FieldExtractor


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Applicant: Acme Water Co")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "water_amount"
    table.cell(0, 1).text = "1000"

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


class FieldExtractorDocumentPipelineTests(unittest.TestCase):
    def test_extract_uses_document_pipeline_for_structured_docx_without_ocr(self) -> None:
        extractor = FieldExtractor()

        with patch.object(extractor, "_download_material", return_value=_docx_bytes()):
            with patch.object(extractor.ocr, "extract_fields") as ocr_extract:
                fields = extractor.extract(
                    MaterialSlot(
                        material_type="APPLICATION_FORM",
                        original_file_name="application.docx",
                        storage_key="materials/application.docx",
                        file_extension="docx",
                        uploaded=True,
                    )
                )

        field_map = {field.field_key: field.field_value for field in fields}

        self.assertIn("content_block_1", field_map)
        self.assertIn("Applicant: Acme Water Co", field_map["content_block_1"])
        self.assertEqual("1000", field_map["water_amount"])
        self.assertTrue(all(field.source_material == "APPLICATION_FORM" for field in fields))
        ocr_extract.assert_not_called()


if __name__ == "__main__":
    unittest.main()
