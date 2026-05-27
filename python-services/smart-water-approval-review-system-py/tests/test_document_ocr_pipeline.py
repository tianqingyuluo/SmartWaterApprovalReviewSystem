from __future__ import annotations

import unittest
from io import BytesIO
from typing import Any

import fitz  # type: ignore[import-untyped]
from docx import Document

from src.models import ExtractedField


def _load_pipeline_class():
    try:
        from src.services.document_ocr_pipeline import DocumentOcrPipeline
    except ModuleNotFoundError as exc:
        raise AssertionError(
            "Expected src.services.document_ocr_pipeline.DocumentOcrPipeline to implement the "
            "production document/OCR parsing pipeline."
        ) from exc
    return DocumentOcrPipeline


def _parse_material(pipeline: Any, *, material_type: str, source_file_name: str, file_bytes: bytes) -> Any:
    for method_name in ("parse_material", "process_material", "parse"):
        method = getattr(pipeline, method_name, None)
        if method:
            return method(
                material_type=material_type,
                source_file_name=source_file_name,
                file_bytes=file_bytes,
            )

    raise AssertionError(
        "Expected document pipeline to expose parse_material(...), process_material(...), or parse(...)"
    )


def _contract_dict(result: Any) -> dict[str, Any]:
    if isinstance(result, dict):
        return result
    if hasattr(result, "model_dump"):
        return result.model_dump(by_alias=True)
    return {
        "materialType": getattr(result, "material_type", None),
        "sourceFileName": getattr(result, "source_file_name", None),
        "contentBlocks": getattr(result, "content_blocks", None),
        "tables": getattr(result, "tables", None),
        "extractedFields": getattr(result, "extracted_fields", None),
        "ocrMetadata": getattr(result, "ocr_metadata", None),
        "errors": getattr(result, "errors", None),
    }


def _field_dict(field: Any) -> dict[str, Any]:
    if isinstance(field, dict):
        return field
    if hasattr(field, "model_dump"):
        return field.model_dump(by_alias=True)
    return {
        "fieldKey": getattr(field, "field_key", None),
        "fieldValue": getattr(field, "field_value", None),
        "confidence": getattr(field, "confidence", None),
        "sourceMaterial": getattr(field, "source_material", None),
    }


def _error_code(error: Any) -> str:
    if isinstance(error, dict):
        return str(error.get("code") or "")
    return str(getattr(error, "code", ""))


def _error_message(error: Any) -> str:
    if isinstance(error, dict):
        return str(error.get("message") or "")
    return str(getattr(error, "message", ""))


def _content_text(block: Any) -> str:
    if isinstance(block, dict):
        return str(block.get("text") or block.get("content") or "")
    return str(getattr(block, "text", getattr(block, "content", "")))


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Applicant: Acme Water Co")
    doc.add_paragraph("Project: North Canal intake")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "field"
    table.cell(0, 1).text = "value"
    table.cell(1, 0).text = "water_amount"
    table.cell(1, 1).text = "1000"

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _text_pdf_bytes() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Water permit application\nApplicant: Acme Water Co")
    data = doc.tobytes()
    doc.close()
    return data


def _blank_pdf_bytes() -> bytes:
    doc = fitz.open()
    doc.new_page()
    data = doc.tobytes()
    doc.close()
    return data


class _StubOcrAdapter:
    def __init__(self, fields: list[ExtractedField]) -> None:
        self.fields = fields
        self.calls: list[tuple[bytes, str, str]] = []

    def extract_fields(self, file_bytes: bytes, material_type: str, file_name: str) -> list[ExtractedField]:
        self.calls.append((file_bytes, material_type, file_name))
        return list(self.fields)


class DocumentOcrPipelineContractTests(unittest.TestCase):
    def _pipeline(self, ocr_adapter: _StubOcrAdapter | None = None):
        pipeline_class = _load_pipeline_class()
        return pipeline_class(ocr_adapter=ocr_adapter or _StubOcrAdapter([]))

    def test_unsupported_file_type_returns_explicit_error_without_calling_ocr(self) -> None:
        ocr_adapter = _StubOcrAdapter([])
        result = _parse_material(
            self._pipeline(ocr_adapter),
            material_type="APPLICATION_FORM",
            source_file_name="application.xlsx",
            file_bytes=b"not supported",
        )

        data = _contract_dict(result)
        self.assertEqual("APPLICATION_FORM", data["materialType"])
        self.assertEqual("application.xlsx", data["sourceFileName"])
        self.assertEqual([], data["contentBlocks"])
        self.assertEqual([], data["tables"])
        self.assertEqual([], data["extractedFields"])
        self.assertIn("UNSUPPORTED_FILE_TYPE", {_error_code(error) for error in data["errors"]})
        self.assertEqual([], ocr_adapter.calls)

    def test_docx_parser_extracts_paragraphs_tables_and_key_fields_without_ocr(self) -> None:
        ocr_adapter = _StubOcrAdapter([])
        result = _parse_material(
            self._pipeline(ocr_adapter),
            material_type="APPLICATION_FORM",
            source_file_name="application.docx",
            file_bytes=_docx_bytes(),
        )

        data = _contract_dict(result)
        block_text = "\n".join(_content_text(block) for block in data["contentBlocks"])
        field_values = {
            _field_dict(field).get("fieldKey"): _field_dict(field).get("fieldValue")
            for field in data["extractedFields"]
        }

        self.assertEqual([], data["errors"])
        self.assertIn("Applicant: Acme Water Co", block_text)
        self.assertGreaterEqual(len(data["tables"]), 1)
        self.assertEqual("1000", field_values.get("water_amount"))
        self.assertEqual([], ocr_adapter.calls)

    def test_text_pdf_parser_extracts_text_without_ocr(self) -> None:
        ocr_adapter = _StubOcrAdapter([])
        result = _parse_material(
            self._pipeline(ocr_adapter),
            material_type="APPLICATION_FORM",
            source_file_name="application.pdf",
            file_bytes=_text_pdf_bytes(),
        )

        data = _contract_dict(result)
        block_text = "\n".join(_content_text(block) for block in data["contentBlocks"])

        self.assertEqual([], data["errors"])
        self.assertIn("Water permit application", block_text)
        self.assertIn("Applicant: Acme Water Co", block_text)
        self.assertEqual([], ocr_adapter.calls)

    def test_parse_failure_records_explicit_error_instead_of_empty_success(self) -> None:
        ocr_adapter = _StubOcrAdapter([])
        result = _parse_material(
            self._pipeline(ocr_adapter),
            material_type="APPLICATION_FORM",
            source_file_name="broken.docx",
            file_bytes=b"this is not a docx file",
        )

        data = _contract_dict(result)

        self.assertEqual([], data["contentBlocks"])
        self.assertEqual([], data["tables"])
        self.assertEqual([], data["extractedFields"])
        self.assertIn("PARSE_ERROR", {_error_code(error) for error in data["errors"]})
        self.assertEqual([], ocr_adapter.calls)

    def test_image_ocr_success_maps_to_unified_content_fields_and_metadata(self) -> None:
        ocr_adapter = _StubOcrAdapter(
            [
                ExtractedField(field_key="company_name", field_value="Acme Water Co", confidence=0.97),
                ExtractedField(field_key="license_no", field_value="LIC-123456", confidence=0.96),
            ]
        )
        result = _parse_material(
            self._pipeline(ocr_adapter),
            material_type="BUSINESS_LICENSE",
            source_file_name="license.jpg",
            file_bytes=b"fake-image-bytes",
        )

        data = _contract_dict(result)
        block_text = "\n".join(_content_text(block) for block in data["contentBlocks"])
        fields = [_field_dict(field) for field in data["extractedFields"]]

        self.assertEqual([], data["errors"])
        self.assertIn("Acme Water Co", block_text)
        self.assertIn({"fieldKey": "license_no", "fieldValue": "LIC-123456"}, [
            {"fieldKey": field.get("fieldKey"), "fieldValue": field.get("fieldValue")} for field in fields
        ])
        self.assertTrue(data["ocrMetadata"]["used"])
        self.assertEqual([(b"fake-image-bytes", "BUSINESS_LICENSE", "license.jpg")], ocr_adapter.calls)

    def test_scanned_pdf_with_no_text_falls_back_to_ocr(self) -> None:
        pdf_bytes = _blank_pdf_bytes()
        ocr_adapter = _StubOcrAdapter(
            [ExtractedField(field_key="ocr_markdown", field_value="Scanned permit text", confidence=0.91)]
        )
        result = _parse_material(
            self._pipeline(ocr_adapter),
            material_type="APPLICATION_FORM",
            source_file_name="scanned.pdf",
            file_bytes=pdf_bytes,
        )

        data = _contract_dict(result)
        block_text = "\n".join(_content_text(block) for block in data["contentBlocks"])

        self.assertEqual([], data["errors"])
        self.assertIn("Scanned permit text", block_text)
        self.assertTrue(data["ocrMetadata"]["used"])
        self.assertEqual([(pdf_bytes, "APPLICATION_FORM", "scanned.pdf")], ocr_adapter.calls)

    def test_ocr_unavailable_records_error_instead_of_fake_success_text(self) -> None:
        ocr_adapter = _StubOcrAdapter(
            [ExtractedField(field_key="ocr_error", field_value="HTTP 503: OCR unavailable", confidence=0.0)]
        )
        result = _parse_material(
            self._pipeline(ocr_adapter),
            material_type="BUSINESS_LICENSE",
            source_file_name="license.png",
            file_bytes=b"fake-image-bytes",
        )

        data = _contract_dict(result)
        block_text = "\n".join(_content_text(block) for block in data["contentBlocks"])
        field_keys = {_field_dict(field).get("fieldKey") for field in data["extractedFields"]}

        self.assertEqual("", block_text)
        self.assertEqual(set(), field_keys - {"ocr_error"})
        self.assertTrue(any("OCR" in _error_code(error) for error in data["errors"]))
        self.assertTrue(any("HTTP 503" in _error_message(error) for error in data["errors"]))
        self.assertTrue(data["ocrMetadata"]["used"])


if __name__ == "__main__":
    unittest.main()
